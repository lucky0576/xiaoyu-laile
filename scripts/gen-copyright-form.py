# -*- coding: utf-8 -*-
"""
把 soft-copyright/软著申请表-填写内容.md 转成 Word，
方便逐项复制到版权中心表单，也方便打印留存。

支持的 Markdown 子集：# 标题、表格、``` 代码块、普通段落、> 引用、列表。
"""
import io
import os
import re

ROOT = r"C:\Users\Administrator\Documents\deepseek-harness\default-workspace\xiaoyu-laile"
SRC = os.path.join(ROOT, "soft-copyright", "软著申请表-填写内容.md")
OUT = os.path.join(ROOT, "soft-copyright", "软著申请表-填写内容.docx")

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

md = io.open(SRC, "r", encoding="utf-8").read()
lines = md.split("\n")

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = Cm(2.4)
sec.bottom_margin = Cm(2.2)
sec.left_margin = Cm(2.6)
sec.right_margin = Cm(2.4)

style = doc.styles["Normal"]
style.font.name = "宋体"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def set_font(run, name="宋体", size=10.5, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def shade(cell, hexcolor="F2F5FA"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


i = 0
while i < len(lines):
    line = lines[i].rstrip()

    # 代码块
    if line.startswith("```"):
        i += 1
        buf = []
        while i < len(lines) and not lines[i].startswith("```"):
            buf.append(lines[i])
            i += 1
        i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.left_indent = Cm(0.6)
        r = p.add_run("\n".join(buf))
        set_font(r, "Consolas", 9)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F7F8FA")
        pPr.append(shd)
        continue

    # 表格
    if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
        header = [c.strip() for c in line.strip().strip("|").split("|")]
        i += 2
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
            i += 1
        table = doc.add_table(rows=1, cols=len(header))
        table.style = "Table Grid"
        for j, h in enumerate(header):
            cell = table.rows[0].cells[j]
            cell.text = ""
            r = cell.paragraphs[0].add_run(h)
            set_font(r, "宋体", 9.5, bold=True)
            shade(cell)
        for row in rows:
            cells = table.add_row().cells
            for j in range(len(header)):
                txt = row[j] if j < len(row) else ""
                txt = txt.replace("<br>", "\n")
                cells[j].text = ""
                for k, seg in enumerate(txt.split("\n")):
                    para = cells[j].paragraphs[0] if k == 0 else cells[j].add_paragraph()
                    r = para.add_run(seg.replace("**", ""))
                    set_font(r, "宋体", 9.5)
        doc.add_paragraph()
        continue

    # 标题
    m = re.match(r"^(#{1,4})\s+(.*)$", line)
    if m:
        level = len(m.group(1))
        text = m.group(2).strip()
        sizes = {1: 16, 2: 13.5, 3: 11.5, 4: 11}
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12 if level <= 2 else 8)
        p.paragraph_format.space_after = Pt(6)
        if level == 1:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_font(r, "黑体", sizes.get(level, 11), bold=True)
        i += 1
        continue

    # 引用
    if line.startswith(">"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        r = p.add_run(line.lstrip("> ").strip())
        set_font(r, "楷体", 10)
        r.font.color.rgb = RGBColor(0x55, 0x5F, 0x6D)
        i += 1
        continue

    # 分隔线
    if line.strip() in ("---", "***"):
        p = doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), "D0D7E2")
        pbdr.append(bottom)
        pPr.append(pbdr)
        i += 1
        continue

    # 普通段落 / 列表
    if line.strip():
        text = line.strip()
        bullet = bool(re.match(r"^(\d+\.|[-*])\s", text))
        text = re.sub(r"^(\d+\.|[-*])\s+", "", text)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        if bullet:
            p.paragraph_format.left_indent = Cm(0.5)
            text = "· " + text
        # 处理 **加粗**
        parts = re.split(r"(\*\*[^*]+\*\*)", text)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                set_font(r, "宋体", 10.5, bold=True)
            else:
                r = p.add_run(part.replace("`", ""))
                set_font(r, "宋体", 10.5)
    else:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)

    i += 1

doc.save(OUT)
print("已生成", os.path.basename(OUT), round(os.path.getsize(OUT) / 1024, 1), "KB")
